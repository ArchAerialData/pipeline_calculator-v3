"""Saved workbooks retain source folders without changing mileage accounting."""
from copy import deepcopy

from openpyxl import load_workbook
import pytest

from pipeline_calculator.core.constants import SURVEY_MILE_METERS
from pipeline_calculator.export.package import export_analysis_package
from pipeline_calculator.export.xlsx import build_analysis_workbook


FOLDER_HEADERS = ["Subfolder", "Folder Path", "Source KML"]


def _record(provenance):
    return {
        "Placemark_ID": "asset-1", "Name": "Asset", "Count": 1,
        "pipelinelength": 1.25, "source_id": 0,
        "interior_meters": SURVEY_MILE_METERS,
        "shared_allocation_meters": SURVEY_MILE_METERS / 4,
        "Shape_Length": 1.25 * SURVEY_MILE_METERS,
        **provenance,
    }


@pytest.mark.parametrize(("provenance", "saved_fields"), [
    ({"folder_name": "Pads", "folder_path": ["Facilities", "Pads"],
      "folder_id": "1/2", "source_kml": "nested/doc.kml"},
     ("Pads", "Facilities / Pads", "nested/doc.kml")),
    ({"folder_name": "", "folder_path": [], "folder_id": "", "source_kml": "doc.kml"},
     ("No subfolder", None, "doc.kml")),
    ({}, ("Folder not recorded", None, None)),
    ({"folder_name": '=SUM(1,2) – 雪', "folder_path": ['=SUM(1,2) – 雪'],
      "folder_id": "0", "source_kml": '=source.kml'},
     ('=SUM(1,2) – 雪', '=SUM(1,2) – 雪', '=source.kml')),
])
def test_saved_point_and_polyline_folders_preserve_text_and_numeric_contract(
        provenance, saved_fields, tmp_path):
    record = _record(provenance)
    results = {"pipelines": [record], "placemarks": [record], "analysis_complete": True}
    snapshot = deepcopy(results)
    destination = tmp_path / "folders.xlsx"
    build_analysis_workbook(results).save(destination)
    workbook = load_workbook(destination)

    pipelines = workbook["Pipeline Length Analysis"]
    points = workbook["Point Pins"]
    assert [cell.value for cell in pipelines[1]][4:] == FOLDER_HEADERS
    assert [cell.value for cell in points[7]][3:] == FOLDER_HEADERS
    assert tuple(cell.value for cell in pipelines[2][4:]) == saved_fields
    assert tuple(cell.value for cell in points[8][3:]) == saved_fields
    for cell in (*pipelines[2][4:], *points[8][3:]):
        if cell.value is not None:
            assert cell.data_type == "s"
    assert pipelines["C2"].value == 1.25
    assert pipelines["C2"].number_format == "0.000"
    assert pipelines["D2"].value == "=SUM(C2:C100000)"
    assert pipelines["D2"].data_type == "f"
    assert points["B2"].value == points["C8"].value == 1
    assert points["C8"].data_type == "n"
    assert points.auto_filter.ref == "A7:F8"
    assert points.freeze_panes == "A8"
    assert results == snapshot


@pytest.fixture
def folder_geography_results():
    provenance = {
        "folder_name": "=Compressor Stations – 北", "folder_id": "0/1",
        "folder_path": ["Facilities", "=Compressor Stations – 北"],
        "source_kml": "=data.kml",
    }
    pipeline = _record(provenance)
    state = {
        "state_code": "TX", "state_name": "Texas", "analysis_complete": True,
        "total_meters": pipeline["Shape_Length"], "total_miles": 1.25,
        "interior_meters": pipeline["interior_meters"],
        "shared_allocation_meters": pipeline["shared_allocation_meters"],
        "interior_savings_meters": 0, "adjusted_total_meters": pipeline["Shape_Length"],
        "pipelines": [pipeline], "overlap_analysis": {"bundled_sections": [], "savings_miles": 0},
    }
    return {
        "pipelines": [pipeline], "placemarks": [pipeline], "analysis_complete": True,
        "total_meters": pipeline["Shape_Length"], "total_miles": 1.25,
        "geography": {
            "status": "complete", "states": [state], "fragments": [{
                "id": "shared-1", "source_id": 0, "source_name": "Asset", "kind": "shared",
                "state_codes": ["TX", "OK"], "length_meters": SURVEY_MILE_METERS / 2,
                **provenance,
            }],
        },
    }


def test_saved_state_tables_and_scoped_workbook_keep_folder_provenance(folder_geography_results, tmp_path):
    destination = tmp_path / "states.xlsx"
    build_analysis_workbook(folder_geography_results).save(destination)
    workbook = load_workbook(destination)
    state_pipelines = workbook["State Pipeline Lengths"]
    shared = workbook["Shared Borders"]
    expected = ("=Compressor Stations – 北", "Facilities / =Compressor Stations – 北", "=data.kml")
    assert [cell.value for cell in state_pipelines[1]][7:] == FOLDER_HEADERS
    assert [cell.value for cell in shared[1]][8:] == FOLDER_HEADERS
    assert tuple(cell.value for cell in state_pipelines[2][7:]) == expected
    assert tuple(cell.value for cell in shared[2][8:]) == expected
    assert [state_pipelines.cell(2, column).value for column in (5, 6, 7)] == [1, 0.25, 1.25]
    assert shared["E2"].value == 0.5
    assert shared["F2"].value == 0.25
    assert state_pipelines.auto_filter.ref == "A1:J2"
    assert shared.auto_filter.ref == "A1:K2"
    for cell in (*state_pipelines[2][7:], *shared[2][8:]):
        assert cell.data_type == "s"

    state = folder_geography_results["geography"]["states"][0]
    scoped_destination = tmp_path / "texas.xlsx"
    build_analysis_workbook(state).save(scoped_destination)
    scoped_workbook = load_workbook(scoped_destination)
    assert "Point Pins" not in scoped_workbook.sheetnames
    assert tuple(cell.value for cell in scoped_workbook["Pipeline Length Analysis"][2][4:]) == expected


def test_export_package_uses_same_folder_columns(folder_geography_results, tmp_path):
    destination = export_analysis_package(folder_geography_results, tmp_path, "folders.kmz", include_maps=False)
    workbook = load_workbook(destination / "analysis.xlsx")
    for name, row, start in [("Pipeline Length Analysis", 1, 4), ("Point Pins", 7, 3),
                             ("State Pipeline Lengths", 1, 7), ("Shared Borders", 1, 8)]:
        assert [cell.value for cell in workbook[name][row]][start:] == FOLDER_HEADERS
