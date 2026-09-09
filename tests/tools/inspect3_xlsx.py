from pathlib import Path
import openpyxl
wb = openpyxl.load_workbook(Path('.claude/Example-XLSX/EXAMPLE_DESIRED-XLSX-EXPORT.xlsx'))
ws2 = wb['Pipeline Overlap Analysis']
cell = ws2.cell(row=2, column=4)
print('R2C4 value:', cell.value)
print('R2C4 data_type:', cell.data_type)
print('R2C4 number_format:', cell.number_format)
print('Formula? ', isinstance(cell.value, str) and cell.value.startswith('='))
frm = None
for r in range(2, min(ws2.max_row, 15)):
    c = ws2.cell(row=r, column=4)
    if isinstance(c.value, str) and c.value.startswith('='):
        frm = (r, c.value)
        break
print('First formula in col4:', frm)
print('Header fills:', [getattr(c.fill.fgColor, 'rgb', None) for c in ws2[1]])
