import sys, json
from pathlib import Path
try:
    import openpyxl
except ImportError:
    print('NO_OPENPYXL')
    sys.exit(0)

p = Path(r".claude/Example-XLSX/EXAMPLE_DESIRED-XLSX-EXPORT.xlsx")
wb = openpyxl.load_workbook(p)
info = {"sheets": []}
for ws in wb.worksheets:
    sheet = {"title": ws.title, "max_row": ws.max_row, "max_column": ws.max_column, "rows": []}
    # capture first 10 rows, first 20 columns of values and styles
    for r in range(1, min(11, ws.max_row+1)):
        row = []
        for c in range(1, min(21, ws.max_column+1)):
            cell = ws.cell(row=r, column=c)
            val = cell.value
            fill = None
            if cell.fill and getattr(cell.fill, 'fgColor', None):
                fill = cell.fill.fgColor.rgb or cell.fill.fgColor.indexed
            font = None
            if cell.font:
                font = {
                    'name': cell.font.name,
                    'size': cell.font.sz,
                    'bold': cell.font.b,
                    'color': getattr(cell.font.color, 'rgb', None) if cell.font.color else None,
                }
            alignment = None
            if cell.alignment:
                alignment = {'horizontal': cell.alignment.horizontal, 'vertical': cell.alignment.vertical}
            row.append({"v": val, "fill": fill, "font": font, "alignment": alignment, "number_format": cell.number_format})
        sheet["rows"].append(row)
    info["sheets"].append(sheet)
print(json.dumps(info, default=str))
