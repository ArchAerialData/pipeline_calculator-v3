import sys
from pathlib import Path
try:
    import openpyxl
except ImportError:
    print('NO_OPENPYXL'); sys.exit(0)
wb = openpyxl.load_workbook(Path('.claude/Example-XLSX/EXAMPLE_DESIRED-XLSX-EXPORT.xlsx'))
ws = wb['Pipeline Length Analysis']
maxc = ws.max_column
print('PLA header:', [ws.cell(row=1, column=c).value for c in range(1, maxc+1)])
print('PLA row2:', [ws.cell(row=2, column=c).value for c in range(1, maxc+1)])
print('PLA col widths:', [ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width for c in range(1, maxc+1)])
print('---')
ws2 = wb['Pipeline Overlap Analysis']
maxc2 = ws2.max_column
print('POA header:', [ws2.cell(row=1, column=c).value for c in range(1, maxc2+1)])
print('POA row2:', [ws2.cell(row=2, column=c).value for c in range(1, maxc2+1)])
print('POA col widths:', [ws2.column_dimensions[openpyxl.utils.get_column_letter(c)].width for c in range(1, maxc2+1)])
