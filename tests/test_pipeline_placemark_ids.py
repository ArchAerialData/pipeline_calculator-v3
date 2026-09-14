"""Source identifiers survive parsing, calculation and spreadsheet round trips."""
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from pipeline_calculator.parsers.kml_kmz import parse_kml_kmz
from pipeline_calculator.core.analyzer import PipelineAnalyzer
from pipeline_calculator.export.xlsx import build_analysis_workbook


@pytest.mark.parametrize('namespace', ['', ' xmlns="http://www.opengis.net/kml/2.2"'])
@pytest.mark.parametrize('extension', ['kml', 'kmz'])
def test_pipeline_placemark_attribute_and_direct_missing_fallback(tmp_path, namespace, extension):
    attributes = ['id="ID_0001"', '', 'id=""', 'id="   "',
                  'other:id="wrong-namespace"', 'id="ID_0001"']
    features = []
    for index, attribute in enumerate(attributes):
        features.append(f'''<Placemark {attribute}><name>Pipeline {index}</name>
          <ExtendedData><Data name="OBJECTID"><value>OBJECT-9</value></Data></ExtendedData>
          <MultiGeometry>
            <LineString id="wrong-line-id"><coordinates>-100,40 -100.001,40</coordinates></LineString>
            <LineString><coordinates>-100,41 -100.001,41</coordinates></LineString>
          </MultiGeometry></Placemark>''')
    source = (f'<kml{namespace} xmlns:other="urn:other"><Document id="wrong-document-id">'
              + ''.join(features) + '</Document></kml>')
    path = tmp_path / f'source.{extension}'
    if extension == 'kmz':
        with ZipFile(path, 'w') as archive:
            archive.writestr('doc.kml', source)
    else:
        path.write_text(source, encoding='utf-8')
    pipelines, _ = parse_kml_kmz(path)
    expected = ['ID_0001', 'N/A', 'N/A', 'N/A', 'N/A', 'ID_0001']
    assert [p['placemark_id'] for p in pipelines] == expected
    # Source IDs do not replace the unique indices used for overlap calculations.
    assert [p['id'] for p in pipelines] == list(range(len(attributes)))
    assert all(len(p['coordinate_paths']) == 2 for p in pipelines)
    analyzer = PipelineAnalyzer()
    rows, total_meters, total_miles = analyzer.calculate_pipeline_lengths(pipelines)
    assert [p['Placemark_ID'] for p in rows] == expected
    workbook = build_analysis_workbook({'pipelines': rows, 'total_meters': total_meters,
                                       'total_miles': total_miles})
    saved_path = tmp_path / 'results.xlsx'
    workbook.save(saved_path)
    saved = load_workbook(saved_path)
    sheet = saved['Pipeline Length Analysis']
    assert sheet['A1'].value == 'Placemark ID'
    assert [sheet.cell(i+2, 1).value for i in range(len(attributes))] == expected
    assert all(sheet.cell(i+2, 1).data_type == 's' for i in range(len(attributes)))
    saved.close()


def test_results_without_placemark_attribute_do_not_reuse_objectid():
    rows, _, _ = PipelineAnalyzer().calculate_pipeline_lengths([
        {'objectid': 'OBJECT-9', 'id': 42, 'name': 'A',
         'coordinates': [(-100, 40), (-100.001, 40)]}])
    assert rows[0]['Placemark_ID'] == 'N/A'
    workbook = build_analysis_workbook({'pipelines': [{'OBJECTID': 'OBJECT-9', 'Name': 'A'}]})
    assert workbook['Pipeline Length Analysis']['A2'].value == 'N/A'
