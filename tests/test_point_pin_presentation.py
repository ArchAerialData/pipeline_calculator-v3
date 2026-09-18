"""Point totals agree across parsed results, the desktop, and saved workbooks."""
from copy import deepcopy

import pytest

from pipeline_calculator.core.analyzer import PipelineAnalyzer
from pipeline_calculator.export.xlsx import build_analysis_workbook


@pytest.fixture
def point_results(tmp_path):
    source = tmp_path / 'mixed.kml'
    source.write_text('''<kml xmlns="http://www.opengis.net/kml/2.2"><Document>
      <Placemark id="mixed"><name>=Mixed()</name>
        <ExtendedData><Data name="OBJECTID"><value>=mixed</value></Data></ExtendedData><MultiGeometry>
        <LineString><coordinates>-100,35 -99.99,35 -99.98,35</coordinates></LineString>
        <Point><coordinates>-100,35</coordinates></Point>
      </MultiGeometry></Placemark>
      <Placemark id="pins"><name>Two pins</name>
        <ExtendedData><Data name="OBJECTID"><value>pins</value></Data></ExtendedData><MultiGeometry>
        <Point><coordinates>-100,35</coordinates></Point>
        <Point><coordinates>-100,35</coordinates></Point>
      </MultiGeometry></Placemark>
    </Document></kml>''', encoding='utf-8')
    return PipelineAnalyzer().analyze_complete(source)


def test_saved_point_sheet_counts_actual_geometries_and_preserves_literals(point_results, tmp_path):
    from openpyxl import load_workbook

    snapshot = deepcopy(point_results)
    workbook = build_analysis_workbook(point_results)
    output = tmp_path / 'points.xlsx'
    workbook.save(output)
    saved = load_workbook(output)
    points = saved['Point Pins']
    assert points['B2'].value == len(point_results['placemarks']) == 3
    assert list(points.iter_rows(min_row=8, values_only=True)) == [
        ('=mixed', '=Mixed()', 1), ('pins', 'Two pins', 1), ('pins', 'Two pins', 1)]
    assert sum(points.cell(row, 3).value for row in range(8, 11)) == points['B2'].value
    assert 'Combined' in points['B3'].value and 'loaded source documents' in points['B3'].value
    assert 'vertices are not pins' in points['B4'].value
    assert points['A8'].data_type == points['B8'].data_type == 's'
    assert saved['Pipeline Length Analysis']['D2'].data_type == 'f'
    assert point_results == snapshot


def test_point_sheet_distinguishes_recorded_zero_and_unrecorded_inventory():
    baseline = {'pipelines': [], 'analysis_complete': True}
    assert 'Point Pins' not in build_analysis_workbook(baseline).sheetnames
    assert 'Point Pins' not in build_analysis_workbook({**baseline, 'placemarks': None}).sheetnames
    sheet = build_analysis_workbook({**baseline, 'placemarks': []})['Point Pins']
    assert sheet['B2'].value == 0
    assert sheet['B5'].value == 'Complete'
    assert list(sheet.iter_rows(min_row=8, values_only=True)) == []
    assert 'Point Pins' not in build_analysis_workbook({**baseline, 'placemarks': [],
                                                       'state_code': 'TX'}).sheetnames


def test_incomplete_run_note_does_not_claim_points_are_state_allocated_or_incomplete(point_results):
    point_results['analysis_complete'] = False
    point_results['diagnostics'] = [{'level': 'error', 'code': 'overlap_analysis_failed',
                                    'message': 'Overlap calculation failed.'}]
    sheet = build_analysis_workbook(point_results)['Point Pins']
    assert sheet['B2'].value == 3
    assert sheet['A5'].value == 'Overall run status'
    assert sheet['B5'].value.startswith('Incomplete; review Diagnostics')
    assert 'Combined' in sheet['B3'].value


@pytest.mark.native_gui
def test_summary_and_point_table_agree_with_workbook_without_state_zero(point_results):
    import customtkinter as ctk
    from tkinter import ttk
    from pipeline_calculator.gui.tabs.placemarks_tab import create
    from pipeline_calculator.gui.tabs.summary_tab import SummaryView
    from test_summary_tab import labels

    def descendants(widget):
        for child in widget.winfo_children():
            yield child
            yield from descendants(child)

    root = ctk.CTk()
    try:
        combined = SummaryView(root, point_results)
        combined.toggle_details()
        assert 'Point pins: 3' in labels(combined.details)
        combined.destroy()

        parent = ctk.CTkFrame(root)
        create(parent, point_results)
        tree = next(widget for widget in descendants(parent) if isinstance(widget, ttk.Treeview))
        assert 'Point pins: 3' in labels(parent)
        table_total = sum(int(tree.item(item, 'values')[2]) for item in tree.get_children())
        assert table_total == build_analysis_workbook(point_results)['Point Pins']['B2'].value == 3
        parent.destroy()

        state = SummaryView(root, {**point_results, 'state_code': 'TX', 'state_name': 'Texas', 'placemarks': []})
        state.toggle_details()
        state_labels = labels(state.details)
        assert 'Point pins: available in the Combined view only.' in state_labels
        assert 'Point pins: 0' not in state_labels and 'Point pins: 3' not in state_labels
        state.destroy()
    finally:
        root.destroy()
