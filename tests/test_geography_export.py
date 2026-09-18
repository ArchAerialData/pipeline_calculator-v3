from __future__ import annotations

from copy import deepcopy
from datetime import datetime
import json
from pathlib import Path
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from pyproj import Geod
import pytest

from pipeline_calculator.core.analyzer import PipelineAnalyzer
from pipeline_calculator.core.constants import SURVEY_MILE_METERS
from pipeline_calculator.export.corridor_kml import build_overlap_corridor_kml
from pipeline_calculator.export.geography_kmz import build_geography_kml
from pipeline_calculator.export.package import export_analysis_package
from pipeline_calculator.export.xlsx import build_analysis_workbook


NS = {"k": "http://www.opengis.net/kml/2.2"}
GEOD = Geod(ellps="GRS80")


@pytest.fixture
def geography_result():
    fragments = []
    specs = [
        ("state", ["OK"], [(-100, 35), (-100, 35.001)]),
        ("state", ["TX"], [(-100, 34.999), (-100, 35)]),
        ("shared", ["OK", "TX"], [(-100, 35), (-99.999, 35)]),
        ("shared", ["OK", "NM"], [(-100, 35.001), (-100, 35.002)]),
        ("outside", [], [(-110, 55), (-110, 55.001)]),
    ]
    for index, (kind, codes, coordinates) in enumerate(specs):
        length = abs(GEOD.inv(*coordinates[0], *coordinates[1])[2])
        fragments.append({
            "id": f"fragment-{index}", "source_id": index, "source_name": "=SAME NAME()",
            "placemark_id": "=USER ID()", "objectid": "duplicate", "source_kml": "doc.kml",
            "path_index": 0, "start_m": 0, "end_m": length, "length_meters": length,
            "coordinates": coordinates, "kind": kind, "state_codes": codes,
        })
    total = sum(fragment["length_meters"] for fragment in fragments)
    states = []
    for code, name in [("TX", "Texas"), ("OK", "Oklahoma"), ("NM", "New Mexico")]:
        owned = [fragment for fragment in fragments if code in fragment["state_codes"]]
        interior = sum(fragment["length_meters"] for fragment in owned if fragment["kind"] == "state")
        allocated = sum(fragment["length_meters"] / len(fragment["state_codes"])
                        for fragment in owned if fragment["kind"] == "shared")
        states.append({
            "state_code": code, "state_name": name, "analysis_complete": True, "status": "complete",
            "interior_meters": interior, "shared_allocation_meters": allocated,
            "total_meters": interior + allocated, "total_miles": (interior + allocated) / SURVEY_MILE_METERS,
            "interior_savings_meters": 0, "adjusted_total_meters": interior + allocated,
            "pipelines": [{"source_id": fragment["source_id"], "Placemark_ID": fragment["placemark_id"],
                           "Name": fragment["source_name"],
                           "interior_meters": fragment["length_meters"] if fragment["kind"] == "state" else 0,
                           "shared_allocation_meters": fragment["length_meters"] / 2 if fragment["kind"] == "shared" else 0,
                           "Shape_Length": fragment["length_meters"] / (2 if fragment["kind"] == "shared" else 1)}
                          for fragment in owned],
            "overlap_analysis": {"bundled_sections": [], "savings_miles": 0}, "diagnostics": [],
        })
    return {
        "pipelines": [{"Placemark_ID": "=USER ID()", "Name": "=SAME NAME()", "pipelinelength": total / SURVEY_MILE_METERS}],
        "total_meters": total, "total_miles": total / SURVEY_MILE_METERS,
        "overlap_analysis": {"bundled_sections": [], "savings_miles": 0},
        "analysis_complete": True, "diagnostics": [],
        "geography": {"schema_version": 1, "status": "complete", "analysis_complete": True,
                      "states": states, "fragments": fragments, "diagnostics": [],
                      "boundary_source": {"vintage": 2025, "url": "https://example.com/boundaries"},
                      "reconciliation": {"combined_meters": total, "difference_meters": 0}},
    }


def test_workbook_geography_fixed_sheets_and_unrounded_allocations(geography_result, tmp_path):
    workbook = build_analysis_workbook(geography_result)
    assert workbook.sheetnames == ["State Summary", "Pipeline Length Analysis", "Pipeline Overlap Analysis",
                                   "State Pipeline Lengths", "State Overlap Analysis", "Shared Borders", "Analysis Details"]
    summary = workbook["State Summary"]
    assert [summary.cell(row, 1).value for row in range(2, 5)] == ["New Mexico", "Oklahoma", "Texas"]
    for row in range(2, 5):
        assert summary.cell(row, 3).value + summary.cell(row, 4).value == pytest.approx(summary.cell(row, 5).value)
    assert workbook["State Pipeline Lengths"].auto_filter.ref is not None
    assert workbook["Shared Borders"].cell(2, 7).value == "Not calculated"
    output = tmp_path / "workbook.xlsx"
    workbook.save(output)
    from openpyxl import load_workbook
    saved = load_workbook(output)
    for sheet in saved:
        for row in sheet:
            for cell in row:
                if cell.value in ("=SAME NAME()", "=USER ID()"):
                    assert cell.data_type == "s"
    assert saved["Pipeline Length Analysis"]["D2"].data_type == "f"


def test_workbook_scoped_failures_and_diagnostics(geography_result):
    state = geography_result["geography"]["states"][0]
    state.update(status="incomplete", analysis_complete=False, adjusted_total_meters=None, interior_savings_meters=None)
    state["diagnostics"] = [{"level": "error", "code": "overlap_analysis_failed", "message": "=failure"}]
    workbook = build_analysis_workbook(geography_result)
    assert workbook["State Summary"].cell(4, 7).value == "Unavailable"
    assert workbook["State Summary"].cell(4, 8).value == "incomplete"
    diagnostics = workbook["Diagnostics"]
    assert diagnostics.cell(2, 3).value == "=failure"
    assert diagnostics.cell(2, 3).data_type == "s"
    assert "Texas" in diagnostics.cell(2, 4).value


def test_state_summary_reconciliation_is_readable_and_outside_filter(geography_result):
    summary = build_analysis_workbook(geography_result)["State Summary"]
    assert summary.auto_filter.ref == "A1:H4"
    facts = {row[0].value: row[4].value for row in summary.iter_rows(min_row=5)}
    assert facts["Combined original mileage (mi)"] == pytest.approx(geography_result["total_miles"])
    assert facts["Outside supported coverage (mi)"] > 0
    assert facts["Reconciliation difference (mi)"] == pytest.approx(0, abs=1e-12)
    assert facts["State breakdown status"] == "Complete"
    geography_result["geography"].update(states=[], status="unavailable")
    summary = build_analysis_workbook(geography_result)["State Summary"]
    assert any(cell.value == "Unavailable" for row in summary for cell in row)
    assert any("No state results" in str(cell.value) for row in summary for cell in row)


def test_package_roundtrip_preserves_combined_and_interior_lengths(geography_result, tmp_path):
    output = export_analysis_package(geography_result, tmp_path, "sample.kmz", include_json=True)
    assert (output / "analysis.xlsx").is_file()
    saved = json.loads((output / "analysis.json").read_text(encoding="utf-8"))
    assert len(saved["geography"]["fragments"]) == 5
    assert not (output / "States/New Mexico").exists()  # Shared-only allocation has no empty map.
    combined = PipelineAnalyzer().analyze_complete(output / "Combined/analysis.kmz")
    assert combined["total_meters"] == pytest.approx(geography_result["total_meters"], abs=0.001)
    for state in geography_result["geography"]["states"]:
        if state["interior_meters"]:
            result = PipelineAnalyzer().analyze_complete(output / f"States/{state['state_name']}/analysis.kmz")
            assert result["total_meters"] == pytest.approx(state["interior_meters"], abs=0.001)
    with ZipFile(output / "Combined/analysis.kmz") as archive:
        assert archive.namelist() == ["doc.kml"]
        tree = ET.fromstring(archive.read("doc.kml"))
    assert len(tree.findall(".//k:LineString", NS)) == 5
    assert "Shared Borders" in [node.text for node in tree.findall(".//k:Folder/k:name", NS)]


def test_package_options_and_collision_safety(geography_result, tmp_path, monkeypatch):
    import pipeline_calculator.export.package as module

    class FrozenDatetime(datetime):
        @classmethod
        def now(cls):
            return cls(2026, 9, 14, 12, 0, 0)

    monkeypatch.setattr(module, "datetime", FrozenDatetime)
    output = export_analysis_package(geography_result, tmp_path, "sample.kml", include_maps=False)
    second = export_analysis_package(geography_result, tmp_path, "sample.kml", include_maps=False)
    assert output.name == "sample_analysis_20260914_120000"
    assert second.name == "sample_analysis_20260914_120000_2"
    assert [path.name for path in output.iterdir()] == ["analysis.xlsx"]
    assert sorted(path.name for path in tmp_path.iterdir()) == [output.name, second.name]


def test_failed_package_is_not_published_and_existing_output_is_untouched(geography_result, tmp_path, monkeypatch):
    import pipeline_calculator.export.package as module

    sentinel = tmp_path / "existing.txt"
    sentinel.write_text("keep")

    def fail(*args, **kwargs):
        raise OSError("disk full")

    monkeypatch.setattr(module, "write_geography_kmz", fail)
    with pytest.raises(OSError, match="disk full"):
        export_analysis_package(geography_result, tmp_path)
    assert list(tmp_path.iterdir()) == [sentinel]
    assert sentinel.read_text() == "keep"


def test_clipped_corridors_keep_holes_and_multipart_without_countable_centerlines(geography_result):
    polygons = [{
        "outer": [(-100, 35), (-99.99, 35), (-99.99, 35.01), (-100, 35.01), (-100, 35)],
        "holes": [[(-99.998, 35.002), (-99.996, 35.002), (-99.996, 35.004), (-99.998, 35.004), (-99.998, 35.002)]],
    }, {"outer": [(-99.98, 35), (-99.97, 35), (-99.97, 35.01), (-99.98, 35)], "holes": []}]
    section = {"pipeline_1": "A & B", "pipeline_2": "C", "clipped_polygons": polygons}
    preview = ET.fromstring(build_overlap_corridor_kml(section, 1))
    assert len(preview.findall(".//k:Polygon", NS)) == 2
    assert len(preview.findall(".//k:innerBoundaryIs", NS)) == 1
    assert not preview.findall(".//k:LineString", NS)
    assert not preview.findall(".//k:Point", NS)
    assert "A &amp; B" in build_overlap_corridor_kml(section, 1)
    state = geography_result["geography"]["states"][0]
    state["overlap_analysis"]["bundled_sections"] = [section]
    map_tree = ET.fromstring(build_geography_kml(geography_result, "TX"))
    assert len(map_tree.findall(".//k:LineString", NS)) == 1
    assert len(map_tree.findall(".//k:Polygon", NS)) == 2


def test_empty_clipped_corridor_never_uses_uncut_fallback(geography_result):
    section = {"clipped_polygons": [], "bbox": {"min_lon": -120, "max_lon": -110, "min_lat": 30, "max_lat": 40}}
    with pytest.raises(ValueError, match="clipped"):
        build_overlap_corridor_kml(section, 1)
    state = geography_result["geography"]["states"][0]
    state["overlap_analysis"]["bundled_sections"] = [section]
    tree = ET.fromstring(build_geography_kml(geography_result, "TX"))
    assert not tree.findall(".//k:Polygon", NS)


def test_map_preserves_tiny_interval_below_display_precision(geography_result, tmp_path):
    result = deepcopy(geography_result)
    fragment = result["geography"]["fragments"][0]
    fragment["coordinates"] = [(-100, 35), (-100, 35.00000001)]
    fragment["length_meters"] = abs(GEOD.inv(*fragment["coordinates"][0], *fragment["coordinates"][1])[2])
    result["geography"]["fragments"] = [fragment]
    file = tmp_path / "tiny.kml"
    file.write_text(build_geography_kml(result), encoding="utf-8")
    parsed = PipelineAnalyzer().analyze_complete(file)
    assert parsed["total_meters"] > 0
    assert parsed["total_meters"] == pytest.approx(fragment["length_meters"], abs=1e-6)


def test_package_json_rejects_non_json_native_objects(geography_result, tmp_path):
    geography_result["unexpected"] = object()
    with pytest.raises(TypeError):
        export_analysis_package(geography_result, tmp_path, include_json=True)
    assert not list(tmp_path.iterdir())


@pytest.mark.native_gui
def test_package_dialog_defaults_export_all_scopes_without_blocking_tk(geography_result, tmp_path, monkeypatch):
    import threading
    import customtkinter as ctk
    from tkinter import filedialog, messagebox
    from pipeline_calculator.gui.actions.export_actions import export_with_dialog
    from pipeline_calculator.gui.window import AppWindow
    import pipeline_calculator.export.package as package

    root = AppWindow()
    root.update()
    captured = {}

    def fake_export(results, directory, filename, **options):
        captured.update(results=results, directory=directory, filename=filename, options=options,
                        worker=threading.current_thread() is not threading.main_thread())
        return tmp_path / "sample_analysis"

    monkeypatch.setattr(package, "export_analysis_package", fake_export)
    monkeypatch.setattr(filedialog, "askdirectory", lambda **kwargs: str(tmp_path))
    monkeypatch.setattr(messagebox, "showinfo", lambda *args, **kwargs: None)
    exceptions = []
    root.report_callback_exception = lambda *args: exceptions.append(args)

    def descendants(widget):
        for child in widget.winfo_children():
            yield child
            yield from descendants(child)

    def interact():
        controls = list(descendants(root))
        checkboxes = [widget for widget in controls if isinstance(widget, ctk.CTkCheckBox)]
        captured["initial_options"] = {widget.cget("text"): widget.get() for widget in checkboxes}
        next(widget for widget in controls if isinstance(widget, ctk.CTkButton)
             and widget.cget("text") == "Choose Folder & Export").invoke()

    try:
        root.after(300, interact)
        exported = export_with_dialog(geography_result, "sample.kmz")
        assert exported == str(tmp_path / "sample_analysis")
        assert not exceptions
        assert captured["worker"]
        assert captured["initial_options"] == {"Include KMZ maps": 1, "Include JSON data": 0}
        assert captured["options"] == {"include_maps": True, "include_json": False}
        assert captured["results"] is geography_result
        assert len(captured["results"]["geography"]["states"]) == 3
    finally:
        root.destroy()
