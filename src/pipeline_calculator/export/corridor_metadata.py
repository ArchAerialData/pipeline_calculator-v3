"""Small, shared presentation and export contract for versioned corridor maps."""
from __future__ import annotations

import math


EXCEL_GEOMETRY_TEXT_LIMIT = 30_000
MAP_NOTE = 'Corridor maps show approximate areas around qualifying paths. Padding does not affect mileage.'
UNAVAILABLE_NOTE = 'Corridor map unavailable; mileage and savings are unchanged. See Diagnostics.'


def is_versioned_corridor(section):
    return 'visualization_schema_version' in section


def corridor_sections(results):
    """Visit all stored scopes, without recomputing or modifying their geometry."""
    yield from (results.get('overlap_analysis') or {}).get('bundled_sections', [])
    for state in (results.get('geography') or {}).get('states', []):
        yield from (state.get('overlap_analysis') or {}).get('bundled_sections', [])


def has_corridor_metadata(results):
    return any(is_versioned_corridor(section) for section in corridor_sections(results))


def corridor_map_status(section):
    """Use the completed decision; never do topology work from a UI callback."""
    if section.get('visualization_status') == 'omitted':
        return 'Unavailable — see Diagnostics'
    polygons = section.get('clipped_polygons', section.get('visualization_polygons'))
    if polygons is not None:
        if not polygons:
            return 'Unavailable — see Diagnostics'
        parts = len(polygons)
        holes = sum(len(polygon.get('holes', [])) for polygon in polygons)
        return f'Available ({parts} {"part" if parts == 1 else "parts"}, {holes} {"hole" if holes == 1 else "holes"})'
    return 'Legacy approximation' if not is_versioned_corridor(section) else 'Unavailable — see Diagnostics'


def corridor_detail_text(results):
    """Summarize actual completed display policies, rather than analysis settings."""
    paddings = set()
    for section in corridor_sections(results):
        if not is_versioned_corridor(section):
            continue
        value = (section.get('visualization_metadata') or {}).get('padding_m')
        if type(value) in (int, float) and math.isfinite(value) and value > 0:
            paddings.add(float(value))
    if len(paddings) == 1:
        return f'Approximate overlap area with {next(iter(paddings)):g} m padding around qualifying paths. Padding does not affect mileage.'
    return MAP_NOTE


def validate_corridor_results(results):
    """Do not let JSON's legacy default=str disguise invalid new result payloads."""
    def native(value, depth=0):
        if depth > 24:
            raise ValueError('Corridor metadata nesting is too deep')
        if value is None or type(value) in (str, bool, int):
            return
        if type(value) is float and math.isfinite(value):
            return
        if type(value) is list:
            for child in value:
                native(child, depth + 1)
            return
        if type(value) is dict and all(type(key) is str for key in value):
            for child in value.values():
                native(child, depth + 1)
            return
        raise ValueError('Corridor results must contain only finite JSON-native values')

    for section in corridor_sections(results):
        if not is_versioned_corridor(section):
            continue
        if type(section['visualization_schema_version']) is not int or section['visualization_schema_version'] != 1:
            raise ValueError('Unsupported corridor visualization schema')
        if section.get('visualization_status') not in ('ready', 'omitted'):
            raise ValueError('Corridor visualization status is invalid')
        for key, value in section.items():
            if key.startswith('visualization_') or key in ('clipped_polygons', 'diagnostics'):
                native(value)
        for key in ('visualization_polygons', 'clipped_polygons'):
            if key not in section:
                continue
            polygons = section[key]
            if type(polygons) is not list:
                raise ValueError('Corridor polygons must be a JSON list')
            for polygon in polygons:
                if type(polygon) is not dict or type(polygon.get('outer')) is not list or type(polygon.get('holes')) is not list:
                    raise ValueError('Corridor polygons require outer and holes lists')
        if section['visualization_status'] == 'omitted':
            if section.get('visualization_polygons') or section.get('clipped_polygons'):
                raise ValueError('An omitted corridor cannot carry available map geometry')
        elif not section.get('clipped_polygons', section.get('visualization_polygons')):
            raise ValueError('A ready corridor must contain its complete canonical polygons')


def bounded_geometry_text(value):
    """Excel silently truncates large cells. Omit the whole optional depiction."""
    if value is None or value == []:
        return ''
    text = str(value)
    return text if len(text) <= EXCEL_GEOMETRY_TEXT_LIMIT else ''


def compatibility_ring_text(section):
    """Only a complete canonical, hole-free single polygon fits the old column."""
    if section.get('visualization_status') != 'ready':
        return ''
    polygons = section.get('clipped_polygons', section.get('visualization_polygons', []))
    if len(polygons) != 1 or polygons[0].get('holes'):
        return ''
    return bounded_geometry_text(polygons[0]['outer'])


def add_corridor_details(workbook, results):
    if not has_corridor_metadata(results):
        return
    from openpyxl.styles import Alignment, Font, PatternFill

    if 'Analysis Details' in workbook:
        sheet = workbook['Analysis Details']
        sheet.append([])
    else:
        sheet = workbook.create_sheet('Analysis Details')
        sheet.append(['Detail', 'Value'])
        sheet.freeze_panes = 'A2'
        sheet.column_dimensions['A'].width = 45
        sheet.column_dimensions['B'].width = 110
        for cell in sheet[1]:
            cell.font = Font(name='Aptos Display', size=11, bold=True, color='FFFFFFFF')
            cell.fill = PatternFill('solid', fgColor='FF285D73')
    first = sheet.max_row + 1
    sheet.append(['Corridor maps', corridor_detail_text(results)])
    policies = sorted({str((section.get('visualization_metadata') or {}).get('policy', ''))
                       for section in corridor_sections(results) if is_versioned_corridor(section)})
    policy_text = ', '.join(policies)
    sheet.append(['Corridor policy', policy_text if len(policy_text) <= 12000 else 'Multiple policies; see JSON metadata.'])
    sheet.append(['Corridor geometry columns',
                  'Legacy center/geometry/width cells are blank when not applicable. Only a complete single polygon without holes may appear in corridor_polygon. '
                  'Geometry exceeding 30,000 characters is omitted from its cell, never truncated; complete shapes are available in corridor maps and optional JSON.'])
    sheet.append(['Corridor map failures', UNAVAILABLE_NOTE])
    for row in sheet.iter_rows(min_row=first):
        for cell in row:
            cell.font = Font(name='Aptos Narrow', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.data_type == 'f':
                cell.data_type = 's'
        sheet.row_dimensions[row[0].row].height = 48
    sheet.auto_filter.ref = sheet.dimensions
