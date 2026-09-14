"""Fixed, filterable geography sheets layered onto the established workbook."""
from __future__ import annotations

import json
import math

from pipeline_calculator.core.constants import SURVEY_MILE_METERS


def _miles(value):
    return None if value is None else float(value) / SURVEY_MILE_METERS


def _text(value):
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def add_geography_sheets(workbook, results):
    """Append geography tables without altering ordinary-mode workbook contracts."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    geography = results["geography"]
    states = sorted(geography.get("states", []), key=lambda s: s["state_name"])

    def table(name, headers, rows, *, numeric=(), widths=None):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append([_text(value) for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="Aptos Display", size=11, bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", fgColor="FF285D73")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for col, header in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = (
                widths[col - 1] if widths else min(42, max(18, len(header) + 2)))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Aptos Narrow", size=11)
                cell.alignment = Alignment(vertical="top")
                if cell.column in numeric:
                    # Keep the underlying number unrounded, while exposing tiny values.
                    cell.number_format = '[>=0.001]0.000;[>0]"<0.001";0.000'
        return sheet

    summary_rows = []
    for state in states:
        summary_rows.append([
            state["state_name"], state["state_code"],
            _miles(state.get("interior_meters", 0)),
            _miles(state.get("shared_allocation_meters", 0)), state.get("total_miles", 0),
            _miles(state["interior_savings_meters"]) if state.get("interior_savings_meters") is not None else "Unavailable",
            _miles(state["adjusted_total_meters"]) if state.get("adjusted_total_meters") is not None else "Unavailable",
            state.get("status", "complete" if state.get("analysis_complete", True) else "incomplete"),
        ])
    summary = table("State Summary", [
        "State", "Code", "Interior mileage (mi)", "Shared allocation (mi)",
        "Original attributed mileage (mi)", "Interior mileage removed (mi)",
        "Adjusted attributed mileage (mi)", "Status",
    ], summary_rows, numeric=(3, 4, 5, 6, 7))
    # The summary is the entry point; retain established combined sheet names.
    workbook.move_sheet(summary, offset=-workbook.index(summary))
    reconciliation = geography.get("reconciliation") or {}
    fragments = geography.get("fragments", [])
    attributed = math.fsum(float(state.get("total_meters", 0)) for state in states)
    outside = reconciliation.get("outside_meters", math.fsum(
        fragment["length_meters"] for fragment in fragments if fragment["kind"] == "outside"))
    unresolved = reconciliation.get("unresolved_meters", math.fsum(
        fragment["length_meters"] for fragment in fragments if fragment["kind"] == "unresolved"))
    combined = float(results.get("total_meters", 0))
    difference = reconciliation.get("delta_meters", attributed + outside + unresolved - combined)
    facts = [
        ("Combined original mileage (mi)", _miles(combined)),
        ("Sum of state attributed mileage (mi)", _miles(attributed)),
        ("Outside supported coverage (mi)", _miles(outside)),
        ("Unresolved mileage (mi)", _miles(unresolved)),
        ("Reconciliation difference (mi)", _miles(difference)),
        ("State breakdown status", geography.get("status", "incomplete").title()),
        ("Reconciliation status", ("Passed" if reconciliation["passed"] else "Failed")
         if "passed" in reconciliation else "Unavailable"),
    ]
    next_row = summary.max_row + 3
    # Facts remain outside the state table's already-established autofilter.
    for label, value in facts:
        summary.merge_cells(start_row=next_row, start_column=1, end_row=next_row, end_column=4)
        summary.cell(next_row, 1, label).font = Font(name="Aptos Narrow", size=11, bold=True)
        cell = summary.cell(next_row, 5, value)
        if isinstance(value, (int, float)):
            cell.number_format = "0.000000000"
        next_row += 1
    notes = [
        "State overlap is calculated independently. State adjusted mileage and mileage removed need not sum to the combined figures.",
        "Shared allocations receive no state overlap discount. State maps contain interior geometry only; shared geometry appears once in the Combined map.",
    ]
    if not states:
        notes.insert(0, "No state results are available. Review the coverage and status above and any Diagnostics sheet for details.")
    for note in notes:
        summary.merge_cells(start_row=next_row, start_column=1, end_row=next_row, end_column=8)
        cell = summary.cell(next_row, 1, note)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(name="Aptos Narrow", size=11)
        summary.row_dimensions[next_row].height = 32
        next_row += 1

    pipeline_rows = []
    overlap_rows = []
    for state in states:
        for pipeline in state.get("pipelines", []):
            pipeline_rows.append([
                state["state_name"], pipeline.get("source_id", ""),
                pipeline.get("Placemark_ID", ""), pipeline.get("Name", ""),
                _miles(pipeline.get("interior_meters", 0)),
                _miles(pipeline.get("shared_allocation_meters", 0)),
                _miles(pipeline.get("Shape_Length", 0)),
            ])
        overlap = state.get("overlap_analysis") or {}
        for section in overlap.get("bundled_sections", []):
            overlap_rows.append([
                state["state_name"], section.get("pipeline_1_id", ""), section.get("pipeline_1", ""),
                section.get("pipeline_2_id", ""), section.get("pipeline_2", ""), section.get("bundled_length_miles", 0),
                section.get("bundled_length_meters", 0), section.get("average_separation", 0),
                section.get("segment_count", 0),
            ])
    table("State Pipeline Lengths", [
        "State", "Source ID", "Placemark ID", "Pipeline name", "Interior mileage (mi)",
        "Shared allocation (mi)", "Original attributed mileage (mi)",
    ], pipeline_rows, numeric=(5, 6, 7))
    table("State Overlap Analysis", [
        "State", "Source 1 ID", "Pipeline 1", "Source 2 ID", "Pipeline 2", "Bundled length (mi)",
        "Bundled length (m)", "Average separation (m)", "Segment count",
    ], overlap_rows, numeric=(6,))

    shared = [fragment for fragment in geography.get("fragments", [])
              if fragment.get("kind") == "shared"]
    if shared:
        names = {state["state_code"]: state["state_name"] for state in states}
        table("Shared Borders", [
            "Fragment ID", "Source ID", "Pipeline name", "Adjoining states",
            "Physical mileage (mi)", "Allocation per state (mi)", "State overlap treatment", "Map location",
        ], ([fragment["id"], fragment["source_id"], fragment.get("source_name", ""),
             ", ".join(names.get(code, code) for code in fragment["state_codes"]),
             _miles(fragment["length_meters"]),
             _miles(fragment["length_meters"] / len(fragment["state_codes"])),
             "Not calculated", "Combined/analysis.kmz — Shared Borders"] for fragment in shared), numeric=(5, 6))

    detail_rows = [
        ["Geography status", geography.get("status", "incomplete")],
        ["Units", "US survey miles; GRS80 geodesic lengths"],
        ["Accounting", "State attributed mileage = interior mileage + shared-border allocation."],
        ["Original-mileage reconciliation", "State attributed + outside coverage + unresolved = combined original mileage."],
        ["Overlap qualification", "Calculated independently within each state; state adjusted mileage and savings need not sum to combined figures."],
        ["Shared-border overlap", "Not calculated. Shared allocations receive no state overlap discount."],
        ["Map mileage", "State maps contain exclusive interior geometry only. Shared geometry appears once in the Combined map."],
        ["Map availability", "KMZ maps are optional; map references apply when maps are included in the export package."],
        ["Point placemarks", "Combined-only; state views analyze pipeline paths."],
        ["Accuracy", "Numerical clipping precision is distinct from source boundary positional accuracy; this is not a surveyed ownership determination."],
        ["Combined original mileage (mi)", results.get("total_miles", 0)],
    ]
    for key, value in (geography.get("boundary_source") or {}).items():
        detail_rows.append([f"Boundary source: {key}", value])
    for key, value in reconciliation.items():
        detail_rows.append([f"Reconciliation: {key}", value])
    for key, value in (results.get("analysis_parameters") or {}).items():
        detail_rows.append([f"Analysis parameter: {key}", value])
    table("Analysis Details", ["Detail", "Value"], detail_rows, widths=(45, 120))

    extra_diagnostics = []
    for diagnostic in geography.get("diagnostics", []):
        extra_diagnostics.append(("Geography", diagnostic))
    for state in states:
        for diagnostic in state.get("diagnostics", []):
            extra_diagnostics.append((state["state_name"], diagnostic))
        for section in (state.get("overlap_analysis") or {}).get("bundled_sections", []):
            for diagnostic in section.get("diagnostics", []):
                extra_diagnostics.append((state["state_name"], diagnostic))
    if extra_diagnostics:
        if "Diagnostics" not in workbook:
            table("Diagnostics", ["Level", "Code", "Message", "Context"], [], widths=(14, 32, 90, 120))
        diagnostics_sheet = workbook["Diagnostics"]
        seen = set()
        for scope, diagnostic in extra_diagnostics:
            if not isinstance(diagnostic, dict):
                diagnostic = {"message": str(diagnostic)}
            identity = (scope, json.dumps(diagnostic, ensure_ascii=False, sort_keys=True))
            if identity in seen:
                continue
            seen.add(identity)
            diagnostics_sheet.append([
                diagnostic.get("level", "warning"), diagnostic.get("code", ""),
                diagnostic.get("message", ""),
                json.dumps({"scope": scope, "details": diagnostic.get("context", {})}, ensure_ascii=False),
            ])
        diagnostics_sheet.auto_filter.ref = diagnostics_sheet.dimensions
    if "Diagnostics" in workbook:
        workbook.move_sheet("Diagnostics", offset=len(workbook.sheetnames) - 1 - workbook.index(workbook["Diagnostics"]))
