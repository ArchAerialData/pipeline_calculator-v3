"""One repair provenance contract for ordinary and geography exports."""
from __future__ import annotations

import json
import math
import re


def validate_input_repair(results):
    """Reject non-JSON report values before any exporter can stringify them."""
    if 'input_repair' not in results:
        return None
    report = results['input_repair']
    if not isinstance(report, dict) or type(report.get('schema_version')) is not int or report['schema_version'] != 1:
        raise ValueError('Unsupported input repair report schema')

    def check(value, depth=0):
        if depth > 32:
            raise ValueError('Input repair report nesting is too deep')
        if value is None or type(value) in (str, bool, int):
            return
        if type(value) is float and math.isfinite(value):
            return
        if type(value) is list:
            for child in value:
                check(child, depth + 1)
            return
        if type(value) is dict and all(type(key) is str for key in value):
            for child in value.values():
                check(child, depth + 1)
            return
        raise ValueError('Input repair provenance must contain only finite JSON-native values')

    check(report)
    # No default=str: handles, dataclasses and geometry objects are not evidence.
    json.dumps(report, allow_nan=False)
    return report


def _cell_text(value):
    text = str(value)
    # Source metadata is literal. Escape illegal XML controls without discarding
    # the evidence or asking openpyxl to write invalid workbook XML.
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', lambda m: f'\\u{ord(m[0]):04x}', text)
    if len(text) > 12000:
        text = text[:12000] + f' … [summary; {len(text):,} characters in full repair report]'
    return text


def add_repair_details(workbook, results):
    """Append to existing Analysis Details, creating it only for repaired runs."""
    report = validate_input_repair(results)
    if report is None:
        return
    from openpyxl.styles import Alignment, Font, PatternFill

    if 'Analysis Details' in workbook:
        sheet = workbook['Analysis Details']
        sheet.append([])
    else:
        sheet = workbook.create_sheet('Analysis Details')
        sheet.append(['Detail', 'Value'])
        sheet.freeze_panes = 'A2'
        sheet.column_dimensions['A'].width = 45
        sheet.column_dimensions['B'].width = 110
        for cell in sheet[1]:
            cell.font = Font(name='Aptos Display', size=11, bold=True, color='FFFFFFFF')
            cell.fill = PatternFill('solid', fgColor='FF285D73')
    first = sheet.max_row + 1
    sheet.append(['Input repair', 'Formatting/import correction only. Source geometry retained; analysis outputs are separate derived files.'])
    for key, value in report.items():
        if isinstance(value, (list, dict)):
            if key == 'rules' and isinstance(value, list) and all(isinstance(v, str) for v in value):
                text = ', '.join(value)
            else:
                text = f'{len(value):,} entries. Full evidence is available in repair Details and JSON export.'
        else:
            text = value
        sheet.append([f'Input repair: {key.replace("_", " ")}', _cell_text(text)])
    for row in sheet.iter_rows(min_row=first):
        for cell in row:
            cell.font = Font(name='Aptos Narrow', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.data_type == 'f':
                cell.data_type = 's'
        sheet.row_dimensions[row[0].row].height = 44
    # This may already include geography settings; update rather than replacing.
    sheet.auto_filter.ref = sheet.dimensions
