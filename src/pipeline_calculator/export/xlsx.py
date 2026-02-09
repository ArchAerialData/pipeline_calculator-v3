from __future__ import annotations


def build_analysis_workbook(current_results):
    """Build an XLSX workbook (openpyxl) from an analysis results dict.

    Kept as a pure function so it can be unit-tested and used by both the legacy
    monolith and the refactored package modules.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as e:  # pragma: no cover
        raise ImportError(
            "openpyxl is required to export .xlsx files. Install with: pip install openpyxl"
        ) from e

    wb = Workbook()

    # ---------------- Pipeline Length Analysis sheet ----------------
    ws = wb.active
    ws.title = "Pipeline Length Analysis"
    ws.freeze_panes = 'A2'

    headers_pla = [
        "Object ID (if available)",
        "Polyline Name (if available)",
        "Pipeline Lengths (US Survey)",
        "TOTAL MILEAGE",
    ]
    ws.append(headers_pla)

    header_font = Font(name="Aptos Display", size=11, bold=True)
    body_font = Font(name="Aptos Narrow", size=11)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    yellow = PatternFill("solid", fgColor="FFFFFF00")
    green = PatternFill("solid", fgColor="FF00B050")
    gray = PatternFill("solid", fgColor="FFD9D9D9")

    for col_idx in range(1, len(headers_pla) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.alignment = center
        cell.fill = gray
    ws.cell(row=1, column=3).fill = yellow
    ws.cell(row=1, column=4).fill = green

    widths = [25.11, 44.89, 29.55, 27.78]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for p in list(current_results.get('pipelines', []) or []):
        obj_id = p.get('OBJECTID') if p.get('OBJECTID') not in (None, "") else "N/A"
        name = p.get('Name', '')
        miles = float(p.get('pipelinelength', 0.0)) if p.get('pipelinelength') is not None else 0.0
        ws.append([obj_id, name, miles, None])

    max_row = ws.max_row
    for r in range(2, max_row + 1):
        ws.cell(row=r, column=1).alignment = center
        ws.cell(row=r, column=1).font = body_font
        ws.cell(row=r, column=2).alignment = left
        ws.cell(row=r, column=2).font = body_font
        c3 = ws.cell(row=r, column=3)
        c3.alignment = center
        c3.font = body_font
        c3.number_format = '0.000'

    ws.cell(row=2, column=4).value = "=SUM(C2:C100000)"
    ws.cell(row=2, column=4).font = Font(name="Aptos Narrow", size=11, bold=True)
    ws.cell(row=2, column=4).alignment = center

    # ---------------- Pipeline Overlap Analysis sheet ----------------
    ws2 = wb.create_sheet("Pipeline Overlap Analysis")
    ws2.freeze_panes = 'A2'

    headers_poa = [
        "Pipeline 1",
        "Pipeline 2",
        "Bundled Length (mi)",
        "TOTAL MILEAGE REMOVED",
        "Bundled Length (m)",
        "Average Separation",
        "Segment Count",
        "Center (Long)",
        "Center (Lat)",
        "bbox",
        "oriented_polygon",
        "oriented_width_m",
        "corridor_polygon",
    ]
    ws2.append(headers_poa)

    for col_idx in range(1, len(headers_poa) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.alignment = center
        cell.fill = gray
    ws2.cell(row=1, column=3).fill = yellow
    ws2.cell(row=1, column=4).fill = green

    widths2 = [44.89, 13.0, 20.33, 28.11, 21.0, 19.11, 20.78, 17.11, 20.0, 107.89, 194.55, 16.44, 255.78]
    for i, w in enumerate(widths2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    def _serialize_bbox(b):
        if not isinstance(b, dict):
            return str(b)
        return f"{{min_lon: {b.get('min_lon')}, max_lon: {b.get('max_lon')}, min_lat: {b.get('min_lat')}, max_lat: {b.get('max_lat')} }}"

    def _serialize_points(seq):
        try:
            if seq is None:
                return ""
            return str(list(seq))
        except Exception:
            return str(seq)

    bundled = []
    try:
        overlap = current_results.get('overlap_analysis')
        if isinstance(overlap, dict):
            bundled = list(overlap.get('bundled_sections', []) or [])
    except Exception:
        bundled = []

    for s in bundled:
        row = [
            s.get('pipeline_1', ''),
            s.get('pipeline_2', ''),
            float(s.get('bundled_length_miles', 0.0) or 0.0),
            None,
            int(round(float(s.get('bundled_length_meters', 0.0) or 0.0))),
            float(s.get('average_separation', 0.0) or 0.0),
            int(s.get('segment_count', 0) or 0),
            float(s.get('center_lon', 0.0) or 0.0),
            float(s.get('center_lat', 0.0) or 0.0),
            _serialize_bbox(s.get('bbox')),
            _serialize_points(s.get('oriented_polygon')),
            float(s.get('oriented_width_m', 0.0) or 0.0),
            _serialize_points(s.get('corridor_polygon')),
        ]
        ws2.append(row)

    max_row2 = ws2.max_row
    for r in range(2, max_row2 + 1):
        for c in (1, 2, 10, 11, 13):
            ws2.cell(row=r, column=c).font = body_font
            ws2.cell(row=r, column=c).alignment = left
        fmt_map = {
            3: '0.000',
            4: '0.000',
            5: '0',
            6: '0.0',
            7: '0',
            8: '0.0000000',
            9: '0.0000000',
            12: '0.0',
        }
        for c, fmt in fmt_map.items():
            cell = ws2.cell(row=r, column=c)
            cell.font = body_font
            cell.alignment = center
            cell.number_format = fmt

    savings = 0.0
    try:
        overlap = current_results.get('overlap_analysis')
        if isinstance(overlap, dict):
            savings = float(overlap.get('savings_miles', 0.0) or 0.0)
    except Exception:
        savings = 0.0
    ws2.cell(row=2, column=4).value = round(savings, 3)
    ws2.cell(row=2, column=4).font = Font(name="Aptos Narrow", size=11, bold=True)
    ws2.cell(row=2, column=4).alignment = center
    ws2.cell(row=2, column=4).number_format = '0.000'

    return wb

