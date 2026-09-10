"""Run synthetic/local or private manifest expectations without inventing ground truth."""
from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from scripts.validation.common import ROOT, digest, environment, metrics, parallel, save_json, write_fixture
from pipeline_calculator.core.analyzer import PipelineAnalyzer
from pipeline_calculator.gui.actions.export_actions import export_results_to_path


def validate_manifest(data):
    if data.get('schema_version') != 1 or not isinstance(data.get('fixtures'), list):
        raise ValueError('Expected manifest schema_version 1 and fixtures list')
    ids = set()
    for item in data['fixtures']:
        if not item.get('id') or item['id'] in ids:
            raise ValueError('Fixture IDs must be present and unique')
        ids.add(item['id'])
        if item.get('status') not in ('analytic', 'unreviewed', 'reviewed'):
            raise ValueError('Unknown expectation status')
        if not item.get('path') or len(item.get('sha256', '')) != 64:
            raise ValueError('Every fixture needs a path and SHA-256')
        if item.get('expectations') and not item.get('basis'):
            raise ValueError('Asserted expectations need an independent basis')
        if item['status'] == 'reviewed' and not item.get('reviewer'):
            raise ValueError('Reviewed fixtures require a reviewer')
        for value in item.get('expectations', {}).values():
            if isinstance(value, dict):
                if not all(k in value for k in ('value', 'absolute_tolerance')):
                    raise ValueError('Numeric expectations require a tolerance')
                if not math.isfinite(value['value']) or not math.isfinite(value['absolute_tolerance']) or value['absolute_tolerance'] < 0:
                    raise ValueError('Numeric expectations must be finite with nonnegative tolerance')
    return data


def local_manifest(output):
    fixtures = []
    for name, paths, expected, kwargs in [
        ('parallel', parallel(), {'pipelines': 2, 'complete': True, 'savings_meters': {'value': 300, 'absolute_tolerance': 5.01}}, {}),
        ('chain', parallel((0, 10, 20)), {'pipelines': 3, 'complete': True, 'savings_meters': {'value': 300, 'absolute_tolerance': 5.01}}, {}),
        ('below_minimum', parallel(length=100), {'complete': True, 'savings_meters': {'value': 0, 'absolute_tolerance': 0}}, {}),
        ('linked', parallel(length=100), {'pipelines': 2, 'complete': True}, {'linked': 1}),
        ('track', parallel(length=100), {'pipelines': 2, 'complete': True}, {'track': True}),
    ]:
        p = write_fixture(output / (name + ('.kmz' if kwargs.get('linked') else '.kml')), paths, **kwargs)
        fixtures.append({'id': name, 'path': str(p.resolve()), 'sha256': digest(p), 'status': 'analytic',
                         'purpose': name, 'parameters': {}, 'expectations': expected,
                         'basis': 'Authored metric fixture; sampled savings allow one 5 m segment at endpoints.'})
    for name, content in [('empty', '<kml/>'), ('malformed', '<kml>'),
                          ('unsupported', '<kml><Placemark><Polygon/></Placemark></kml>')]:
        p = output / f'{name}.kml'
        p.write_text(content, encoding='utf-8')
        fixtures.append({'id': name, 'path': str(p.resolve()), 'sha256': digest(p), 'status': 'analytic',
                         'basis': 'Explicit empty/unsupported/malformed input',
                         'expectations': {} if name == 'malformed' else {'complete': False},
                         'expected_error': name == 'malformed'})
    fixtures.append({'id': 'tracked_centerlines', 'path': str(ROOT / '.danny/Centerlines.kmz'),
                     'sha256': '2f240c0b680c0afa211e6f0d6d37c1a8ce3bd7386dc00ecbf0166f68a02e7f4f',
                     'status': 'unreviewed', 'expectations': {}, 'purpose': 'Local regression characterization; provenance unknown'})
    return {'schema_version': 1, 'fixtures': fixtures}


def run_manifest(data, output, private_root=None, strict=False):
    validate_manifest(data)
    rows = []
    for item in data['fixtures']:
        path = Path(item['path'])
        if not path.is_absolute():
            path = Path(private_root or ROOT) / path
        row = {'id': item['id'], 'expectation_status': item['status'], 'sha256': item['sha256']}
        rows.append(row)
        if not path.exists():
            row['status'] = 'failed' if strict or item.get('required') else 'missing-private-input'
            continue
        if digest(path) != item['sha256']:
            row.update(status='failed', reason='hash mismatch')
            continue
        try:
            result = PipelineAnalyzer(**item.get('parameters', {})).analyze_complete(str(path))
        except ValueError:
            row['status'] = 'passed' if item.get('expected_error') else 'failed'
            row['error'] = 'analysis raised ValueError'  # No private paths/names in shared summaries.
            continue
        actual = metrics(result)
        row['actual'] = actual
        failures = []
        if item.get('expected_error'):
            failures.append('expected analysis error')
        for key, expected in item.get('expectations', {}).items():
            value = actual.get(key)
            if isinstance(expected, dict):
                ok = isinstance(value, (int, float)) and abs(value - expected['value']) <= expected['absolute_tolerance']
            else:
                ok = value == expected
            if not ok:
                failures.append(key)
        row['status'] = 'failed' if failures else ('unreviewed' if item['status'] == 'unreviewed' else 'passed')
        row['failed_fields'] = failures
        if item['status'] == 'analytic':
            # Export only synthetic data; IDs are not used as filesystem paths.
            stem = output / f'export-{len(rows):03d}'
            export_results_to_path(result, str(stem.with_suffix('.json')))
            export_results_to_path(result, str(stem.with_suffix('.xlsx')))
            from openpyxl import load_workbook
            wb = load_workbook(stem.with_suffix('.xlsx'))
            assert wb.sheetnames
            wb.close()
            assert json.loads(stem.with_suffix('.json').read_text())['total_meters'] == result['total_meters']
    return rows


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--suite', choices=['local'], default='local')
    p.add_argument('--manifest', type=Path)
    p.add_argument('--private-root', type=Path)
    p.add_argument('--strict', action='store_true')
    p.add_argument('--output', type=Path, required=True)
    args = p.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    data = json.loads(args.manifest.read_text()) if args.manifest else local_manifest(args.output)
    rows = run_manifest(data, args.output, args.private_root, args.strict)
    save_json(args.output / 'results.json', {'environment': environment(), 'fixtures': rows})
    (args.output / 'report.md').write_text('# Corpus results\n\n' + '\n'.join(
        f"- {r['id']}: {r['status']} ({r['expectation_status']})" for r in rows) + '\n', encoding='utf-8')
    print(json.dumps({r['id']: r['status'] for r in rows}))
    return int(any(r['status'] == 'failed' for r in rows))


if __name__ == '__main__':
    raise SystemExit(main())
